"""Trace selected frozen numerical transcriptions to their frozen primary sources."""

import argparse
import csv
import hashlib
import html
import json
import re
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

a = argparse.ArgumentParser(description=__doc__)
a.add_argument("--root", type=Path, required=True, help="Frozen project input root.")
a.add_argument(
    "--out", type=Path, required=True, help="Directory for independent source-audit artifacts."
)
args = a.parse_args()
R, W = args.root.resolve(), args.out.resolve()
W.mkdir(parents=True, exist_ok=True)
out = {"sha256": {}, "workbooks": {}, "pdf": {}}
previous = (
    json.loads((W / "source_trace.json").read_text(encoding="utf-8"))
    if (W / "source_trace.json").exists()
    else {"sha256": {}, "pdf": {}}
)

for p in ("data/raw/cochilco/production_1996_2025.xlsx", "data/raw/usgs/myb1-2023-cobalt.xlsx"):
    out["sha256"][p] = hashlib.sha256((R / p).read_bytes()).hexdigest()
    x = load_workbook(R / p, data_only=True, read_only=True)
    print(p, x.sheetnames)
    if "cochilco" in p:
        out["workbooks"][p] = {
            k: x["Hoja1"][k].value for k in ("AD20", "AD21", "AD22", "AD23", "AD24")
        }
    else:
        out["workbooks"][p] = {
            s: [[x[s].cell(i, j).value for j in (1, 2, 3, 11)] for i in range(1, 25)]
            for s in x.sheetnames
            if s == "T8"
        }
    print(out["workbooks"][p])

for p in (
    "data/events/raw/oecd-report-2026.pdf",
    "data/events/raw/oecd-methodology.pdf",
    "data/events/raw/zimbabwe-si213-2022.pdf",
    "data/raw/xu2020/supplement.pdf",
):
    out["sha256"][p] = hashlib.sha256((R / p).read_bytes()).hexdigest()
    if previous["sha256"].get(p) == out["sha256"][p] and p in previous["pdf"]:
        out["pdf"][p] = previous["pdf"][p]
        continue
    x = PdfReader(R / p)
    if "methodology" in p:
        I = list(range(10, min(38, len(x.pages))))
    elif "oecd" in p:
        I = [i for i, z in enumerate(x.pages) if "Table 4.2" in z.extract_text()]
        I = sorted({j for i in I for j in range(max(0, i - 1), min(len(x.pages), i + 3))})
    elif "xu2020" in p:
        I = [
            i
            for i, z in enumerate(x.pages)
            if "Table 23" in z.extract_text()
            or "Table 4" in z.extract_text()
            or "Table 1." in z.extract_text()
        ]
    else:
        I = list(range(len(x.pages)))
    out["pdf"][p] = {str(i + 1): x.pages[i].extract_text() for i in I}
    for i in I:
        (W / (Path(p).stem + f"-p{i + 1}.txt")).write_text(
            x.pages[i].extract_text(), encoding="utf-8"
        )
        if "methodology" in p:
            (W / (Path(p).stem + f"-p{i + 1}-layout.txt")).write_text(
                x.pages[i].extract_text(extraction_mode="layout"), encoding="utf-8"
            )
    print(p, "selected physical pages", [i + 1 for i in I])

Z = list(csv.DictReader((R / "data/raw/baci/country_codes_V202601.csv").open(encoding="utf-8-sig")))
iso = {r["country_name"].casefold(): r["country_iso3"] for r in Z}
iso.update(
    {
        "china (people's republic of)": "CHN",
        "democratic republic of the": "COD",
        "dem rep of congo": "COD",
        "russia": "RUS",
        "türkiye": "TUR",
        "korea": "KOR",
        "russian federation": "RUS",
        "viet nam": "VNM",
        "united states": "USA",
    }
)
selected = {
    "Aluminium",
    "Cobalt",
    "Copper",
    "Graphite",
    "Lithium",
    "Manganese",
    "Nickel",
    "Silicon",
}
covered, countries, unresolved = [], set(), []
country = None
for physical in range(11, 38):
    lines = (
        (W / f"oecd-methodology-p{physical}-layout.txt").read_text(encoding="utf-8").splitlines()
    )
    header = next(line for line in lines if "Product researched" in line)
    pcol = header.index("Product researched")
    yc = [(int(m[0]), (m.start() + m.end()) / 2) for m in re.finditer(r"20\d\d", header)]
    for line in lines:
        marks = list(re.finditer(r"\b(?:vn|x)\b", line))
        if not marks or marks[0].start() < yc[0][1] - 5:
            continue
        c = line[:pcol].strip().replace(" (EU)", "").replace("**", "")
        if c:
            country = c
            countries.add(c)
        p = line[pcol : marks[0].start()].strip()
        if p not in selected:
            continue
        code = iso.get(country.casefold())
        if code is None:
            unresolved.append({"page": physical, "country": country, "product": p})
            continue
        for m in marks:
            y, pos = min(yc, key=lambda z: abs(z[1] - (m.start() + m.end()) / 2))
            if 2012 <= y <= 2024:
                covered.append(
                    {
                        "country": code,
                        "commodity": p,
                        "year": str(y),
                        "surveyed": "1",
                        "restriction_reported": str(int(m[0] == "x")),
                        "source_marker": m[0],
                        "source_page": str(physical),
                    }
                )
F = list(csv.DictReader((R / "data/events/oecd_coverage.csv").open(encoding="utf-8-sig")))
key = lambda r: (r["country"], r["commodity"], r["year"])
A, B = {key(r): r for r in covered}, {key(r): r for r in F}
out["coverage"] = {
    "PDF_rows": len(covered),
    "frozen_rows": len(F),
    "country_labels": sorted(countries),
    "unresolved_country_labels": unresolved,
    "PDF_only": [A[k] for k in sorted(A.keys() - B.keys())],
    "frozen_only": [B[k] for k in sorted(B.keys() - A.keys())],
    "cell_differences": [
        {"PDF": A[k], "frozen": B[k]} for k in sorted(A.keys() & B.keys()) if A[k] != B[k]
    ],
}
if not unresolved:
    with (W / "coverage_from_pdf.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, F[0].keys())
        w.writeheader()
        w.writerows(A[k] for k in sorted(A))
print("coverage", {k: len(v) if isinstance(v, list) else v for k, v in out["coverage"].items()})
out["national_act_raw_excerpts"] = {}
for p, terms in (
    ("data/events/raw/mofcom-2023-39.html", ("2023 年 12 月 1 日", "天然鳞片石墨", "人造石墨")),
    ("data/events/raw/mofcom-2024-46-original.html", ("2024年12月3日", "石墨", "美国")),
    (
        "data/events/raw/indonesia-2019-ministry-release.html",
        ("1 Januari 2020", "1,7", "2 September 2019"),
    ),
):
    data = (R / p).read_bytes()
    out["sha256"][p] = hashlib.sha256(data).hexdigest()
    text = re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", data.decode("utf-8"))))
    excerpts = {}
    for term in terms:
        j = text.casefold().find(term.casefold())
        excerpts[term] = text[max(0, j - 90) : j + len(term) + 180] if j >= 0 else None
    out["national_act_raw_excerpts"][p] = excerpts
    (W / (Path(p).stem + "-plain.txt")).write_text(text + "\n", encoding="utf-8")
    print("national", p, {str(i): excerpts[t] is not None for i, t in enumerate(terms)})
(W / "source_trace.json").write_text(
    json.dumps(out, indent=2, ensure_ascii=False, default=str) + "\n", encoding="utf-8"
)
