"""Extract W5 facts from locked primary bytes: python -m data.research.supplement.

METI charge output is not plant capacity. Eurostat PRODQNT is sold quantity.
Comtrade row sums are diagnostics with unverified overlap/completeness, not production.
"""
import csv
import io
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
EU = {'AT', 'BE', 'BG', 'HR', 'CY', 'CZ', 'DK', 'EE', 'FI', 'FR', 'DE', 'GR', 'HU', 'IE',
      'IT', 'LV', 'LT', 'LU', 'MT', 'NL', 'PL', 'PT', 'RO', 'SK', 'SI', 'ES', 'SE', 'EU27_2020'}


def frozen(path, data):
    """Write a new numerical transcription; existing frozen bytes must agree."""
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        assert path.read_bytes() == data, path
    else:
        path.write_bytes(data)


def main():
    raw = ROOT / 'data/raw'
    W = load_workbook(raw / 'meti/h2daa2024k.xlsx', read_only=True, data_only=True)
    s = W['実数表']
    rows = list(s.iter_rows(min_row=3488, max_row=3498, values_only=True))
    out = io.StringIO(newline='')
    writer = csv.writer(out, lineterminator='\n')
    writer.writerow(['item', 'row', 'month', 'production_kAh'])
    total = 0
    for item, k in [('0108', 3488), ('0109', 3498)]:
        r = rows[k-3488]
        assert str(r[1]).zfill(4) == item
        for month, q in enumerate(r[8:20], 1):
            assert isinstance(q, (int, float)) and q >= 0
            writer.writerow([item, k, f'2024-{month:02}', q])
            total += q
    frozen(raw / 'meti/liion_2024.csv', out.getvalue().encode())

    eurostat = {}
    for code in ['20121955', '20134351', '20134352', '201343Z1', '20134281']:
        rows = list(csv.DictReader((raw / f'eurostat/unfiltered/prodcom-{code}.csv')
                                   .open(encoding='utf-8-sig')))
        selected = [x for x in rows if x['reporter'] in EU and
                    x['indicators'] in {'PRODQNT', 'PQNTFLAG', 'PQNTBASE', 'APRODQNT',
                                        'APQNTFLAG', 'QNTUNIT', 'PRODVAL', 'PVFLAG'}]
        out = io.StringIO(newline='')
        writer = csv.DictWriter(out, fieldnames=list(rows[0]), lineterminator='\n')
        writer.writeheader()
        writer.writerows(selected)
        frozen(raw / f'eurostat/prodcom-{code}-EU.csv', out.getvalue().encode())
        eurostat[code] = {c: {x['indicators']: x['OBS_VALUE'] for x in selected
                             if x['reporter'] == c} for c in sorted(EU)}
        assert not any(x['indicators'] == 'APRODQNT' for x in selected)

    baci = list(csv.DictReader((raw / 'baci/BACI_HS22_Y2024_V202601_selected.csv').open()))
    selected_codes = {r['k'] for r in baci}
    chunks = json.loads((raw / 'comtrade/chunks.json').read_text())
    tariffs = []
    for code in ['28429040', '28539030', '28539050', '2504101000', '2504109100',
                 '3801100030', '3824999940']:
        x = json.loads((raw / f'comtrade/{code}.json').read_text())
        rows = x['data']
        assert x['count'] == len(rows) < 500
        assert all(r['cmdCode'] == code and r['reporterCode'] == 156 and
                   r['period'] == 2024 and r['flowCode'] == 'X' for r in rows)
        dims = ('period', 'reporterCode', 'flowCode', 'partnerCode', 'partner2Code',
                'classificationCode', 'cmdCode', 'customsCode', 'mosCode', 'motCode')
        counts = Counter(tuple(r[k] for k in dims) for r in rows)
        q = sum(float(r['q']) for r in baci if r['i'] == '156' and r['k'] == code[:6]
                and r['q'])  # BACI q is tonnes; preview netWgt is kg.
        if code[:6] not in selected_codes:
            q = None
        z = sum(r['netWgt'] or 0 for r in rows)/1000
        check = next((a for a in chunks if a['code'] == code), None)
        tariffs.append({'code': code, 'rows': len(rows),
                        'partners': len({r['partnerCode'] for r in rows}),
                        'duplicate_reporting_keys': sum(n-1 for n in counts.values()),
                        'returned_row_mass_t': z, 'BACI_HS22_parent': code[:6],
                        'BACI_parent_export_t': q,
                        'diagnostic_row_sum_to_parent': z/q if rows and q else None,
                        'partner_chunks_match': check['known_partner_chunks_match'] if check
                        else None,
                        'status': 'overlap_and_completeness_unverified' if rows else
                        'no_rows_returned_not_evidence_of_zero_trade'})
    result = {'period': 2024, 'japan_kAh': total, 'eurostat': eurostat,
              'eurostat_basis': 'sold quantity; no APRODQNT observations; flags retained verbatim',
              'eurostat_attribution': 'Source: Eurostat DS-059358, extracted 2026-10-16; '
                                      'EU production indicators selected. Eurostat is not '
                                      'responsible for these transformations or conclusions.',
              'comtrade': tariffs, 'comtrade_chunks': chunks,
              'comtrade_basis': 'trade; row-sum diagnostic only; no production replacement',
              'korea': 'No qualifying domestic physical cell output located in public KOSIS/KBIA '
                       'search; KBIA industry-stat pages redirected repeatedly. '
                       'Trade basis retained.',
              'EU_adoption': 'German combined-purity sold carbonate is an imputed variant; '
                             'LFP-only quantity cannot replace aggregate CAM; '
                             'confidential is missing.'}
    frozen(raw / 'supplement/reconciliation.json',
           (json.dumps(result, indent=1, ensure_ascii=True)+'\n').encode())
    print('24 METI monthly observations; 5 EU slices;', len(tariffs), 'tariff lines')


if __name__ == '__main__':
    main()
